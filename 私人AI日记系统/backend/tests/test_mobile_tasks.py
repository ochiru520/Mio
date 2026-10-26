"""Phone task integration with synthetic files, fake model calls and fake ACKs."""
import asyncio
from contextlib import closing, ExitStack
from dataclasses import replace
import io
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document

from app import db, mobile_task_service as mobile, mobile_file_service as files
from app import agent_task_service as tasks, agent_file_service, artifact_service, wechat_service as wx
from app.config import settings
from app.agent_handoff_service import create_handoff
from app.agent_tool_service import ToolExecutionContext
from app.routes import onebot
from app import onebot_transport_service
from tests.test_agent_model_first import _completion
from app.llm import ToolCall
from test_runtime_safety_v2 import isolated

OWNER = {'self_id': '90001', 'user_id': '10001'}
OTHER = {'self_id': '90001', 'user_id': '10002'}


@pytest.fixture(autouse=True)
def phone_setup(isolated, monkeypatch):
    object.__setattr__(settings, 'qq_allowed_user_ids', ('10001', '10002'))
    object.__setattr__(settings, 'qq_bot_enabled', True)
    monkeypatch.setattr(mobile, '_locks', {})
    monkeypatch.setattr(files, '_intake_locks', {})
    monkeypatch.setattr(onebot, '_mobile_intake_locks', {})
    connections = {}
    monkeypatch.setattr(onebot, '_active_connections', connections)
    monkeypatch.setattr(onebot_transport_service, '_active_connections', connections)
    monkeypatch.setattr(wx, '_state', {'connected': False, 'message': '尚未连接微信'})
    monkeypatch.setattr(wx, '_inflight', set())
    monkeypatch.setattr(wx, '_stopping', set())
    monkeypatch.setattr(wx, '_login', {})
    mobile.initialize()


def event(message='/任务 新建 整理报告', *, user='10001', identifier=1, bot='90001'):
    return {'post_type':'message','message_type':'private','self_id':bot,'user_id':user,
            'message_id':identifier,'message':message}


def handoff(origin=OWNER, key='test-handoff'):
    target, _ = mobile.input_context(origin)
    context = ToolExecutionContext(run_id=key, request_id=key, trace_id=key,
        conversation_id=mobile.conversation_id(origin), source=mobile.channel(origin),
        user_message='整理报告', source_message_id=1, step_index=0)
    with patch('app.agent_handoff_service.model_selection', return_value=('test-model', 'low')):
        result = create_handoff({'goal':'整理报告'}, context, delivery_target=target)
    return tasks.get(result['task_id'])


def outbox():
    with closing(db.get_conn()) as conn:
        return [dict(row) for row in conn.execute('SELECT * FROM mobile_task_outbox ORDER BY id')]


def test_duplicate_inbound_creates_one_task_and_one_receipt():
    async def run():
        with patch('app.agent_handoff_service.model_selection', return_value=('test-model', 'low')):
            assert await onebot.process_onebot_event(event(), AsyncMock())
            assert await onebot.process_onebot_event(event(), AsyncMock())
        assert len(mobile.owner_tasks(OWNER)) == 1
        assert len(outbox()) == 1
    asyncio.run(run())


@pytest.mark.parametrize('source', ['qq', 'wechat'])
def test_natural_handoff_and_followup_use_original_instruction(source):
    async def run():
        from app import chat_service as chat
        origin = OWNER
        if source == 'wechat':
            value = {'enabled':True,'token':'synthetic','bot_id':'test@im.bot','user_id':'owner','base_url':wx.API}
            wx.save_config(value); origin = wx.origin_for(value)
        conversation = mobile.conversation_id(origin)
        completion = _completion(ToolCall('new','handoff_to_agent',json.dumps({'goal':'整理文件'})))
        model = AsyncMock(return_value=(completion,'low',''))
        with ExitStack() as stack:
            mocks = {
                'resolve_model_id':'test-model', 'require_configured':None,
                'get_model_profile':SimpleNamespace(model='test-model',supports_tool_calls=True),
                'normalize_model_reasoning':'low', 'load_manuals':{}, 'build_system_prompt':'test',
                '_build_current_time_context':'', 'system_audio_service.chat_context':'',
            }
            for name,value in mocks.items(): stack.enter_context(patch('app.chat_service.'+name,return_value=value))
            stack.enter_context(patch('app.chat_service.build_chat_context',AsyncMock(return_value=SimpleNamespace(system_context='',raw_messages=[]))))
            stack.enter_context(patch('app.chat_service.perform_web_lookup',AsyncMock(return_value=None)))
            stack.enter_context(patch('app.chat_service._self_snapshot_context_for_message',AsyncMock(return_value='')))
            stack.enter_context(patch('app.chat_service._complete_chat_reply_with_single_fallback',model))
            stack.enter_context(patch('app.companion_service.set_pet_activity'))
            stack.enter_context(patch('app.companion_service.infer_speech_emotion',return_value='neutral'))
            stack.enter_context(patch('app.agent_handoff_service.model_selection',return_value=('test-model','low')))
            result = await chat.chat_with_ai('帮我整理文件',conversation_id=conversation,source=source,
                request_id='native-new-'+source,agent_tools_enabled=False,handoff_enabled=True,task_origin=origin)
            task_id = result.task_handoff['task_id']
            model.return_value=(_completion(ToolCall('follow','follow_up_mobile_task',json.dumps({'task_id':task_id}))),'low','')
            follow = await chat.chat_with_ai('控制在两页，别删原始材料',conversation_id=conversation,source=source,
                request_id='native-follow-'+source,agent_tools_enabled=False,handoff_enabled=True,task_origin=origin)
        assert follow.task_handoff['task_id'] == task_id
        assert tasks.get(task_id)['snapshot']['updates'][-1] == '控制在两页，别删原始材料'
        assert len(mobile.owner_tasks(origin)) == 1
        assert {item['function']['name'] for item in model.await_args.kwargs['tools']} == {'handoff_to_agent','follow_up_mobile_task'}
    asyncio.run(run())


def test_confirmation_is_owner_bound_single_use_and_invalidated_by_amendment():
    async def pending(task):
        from app.agent_loop_service import run_agent_loop
        model=AsyncMock(return_value=_completion(ToolCall('profile','update_profile','{"instruction":"以后简洁回答"}')))
        with patch('app.agent_loop_service.call_chat_completion_result',model), patch('app.agent_tool_service.tool_availability',return_value=(True,'')):
            await run_agent_loop(conversation_id=task['conversation_id'],source='agent_background',
                user_message='修改回答方式',source_message_id=1,request_id='confirm-'+task['id'],trace_id='test',
                model_id='test-model',reasoning_level='low',task_id=task['id'])
        current=tasks.get(task['id'])
        assert current['status']=='waiting_confirmation'
        assert '简洁回答' in mobile.action_prompts(current)[0]
        with closing(db.get_conn()) as conn:
            return conn.execute('SELECT token FROM mobile_task_confirmations WHERE task_id=?',(task['id'],)).fetchone()[0]
    async def run():
        task=handoff(); token=await pending(task)
        with pytest.raises(ValueError): await mobile.confirm(OTHER,token,accept=True)
        approve=AsyncMock(return_value={'status':'done'})
        with patch('app.companion_action_service.approve_companion_action',approve):
            await mobile.confirm(OWNER,token,accept=True)
            with pytest.raises(ValueError): await mobile.confirm(OWNER,token,accept=True)
        assert approve.await_count==1
        second=handoff(OWNER,'second'); token=await pending(second)
        await mobile.amend(OWNER,second['id'],'不要修改属性，只回答问题')
        with pytest.raises(ValueError): await mobile.confirm(OWNER,token,accept=True)
    asyncio.run(run())


def test_unconnected_owner_backlog_does_not_starve_online_owner():
    async def run():
        for i in range(25): mobile.queue(OWNER,'offline',key='offline-'+str(i))
        online={**OTHER,'self_id':'90002'}
        mobile.queue(online,'online',key='online')
        api=AsyncMock(return_value={'status':'ok','retcode':0})
        with patch.object(onebot_transport_service,'has_task_connection',side_effect=lambda identifier:identifier=='90002'), patch.object(onebot_transport_service,'request_task_api',api):
            assert await mobile.flush_outbox()==1
        assert outbox()[-1]['status']=='sent'
        assert all(row['status']=='pending' for row in outbox()[:-1])
    asyncio.run(run())


def test_task_ownership_includes_bot_and_sender():
    task = handoff()
    assert not mobile.owner_tasks(OTHER)
    for origin in [OTHER, {**OWNER, 'self_id':'90002'}]:
        with pytest.raises(ValueError): mobile.owned_task(origin, task['id'])
    object.__setattr__(settings, 'qq_allowed_user_ids', ())
    assert not mobile.owner_tasks(OWNER)


def test_handoff_rejects_spoofed_conversation():
    context = ToolExecutionContext(run_id='x', request_id='x', trace_id='', conversation_id='desktop_chat',
        source='qq', source_message_id=1, user_message='test', step_index=0)
    with pytest.raises(ValueError): create_handoff({'goal':'test'}, context, delivery_target=OWNER)


def test_docx_input_is_bound_to_exact_task_and_hash():
    async def run():
        document = Document(); document.add_paragraph('仅用于测试的会议材料。')
        buffer = io.BytesIO(); document.save(buffer)
        await files.store_document(OWNER, 'file-1', '会议.docx', buffer.getvalue())
        target, _ = mobile.input_context(OWNER)
        assert len(target['attachment_ids']) == 1
        task = handoff()
        with closing(db.get_conn()) as conn:
            path = conn.execute('SELECT path FROM mobile_task_attachments').fetchone()[0]
        assert '测试' in agent_file_service.read_document(path, task_id=task['id'])['text']
        with pytest.raises(ValueError): agent_file_service.read_document(path)
        with pytest.raises(ValueError): agent_file_service.read_document(path, task_id=handoff(OTHER, 'other')['id'])
        from pathlib import Path
        Path(path).write_bytes(b'tampered')
        with pytest.raises(ValueError): agent_file_service.read_document(path, task_id=task['id'])
    asyncio.run(run())


@pytest.mark.parametrize('name', ['../secret.txt', 'CON.txt', 'a:stream.txt', 'a.exe', '.hidden.txt', 'bad?.docx'])
def test_unsafe_attachment_names_rejected(name):
    with pytest.raises(ValueError): files.clean_name(name)


def test_corrupt_document_never_becomes_task_input():
    async def run():
        with pytest.raises(Exception): await files.store_document(OWNER, 'file-2', 'broken.docx', b'bad zip')
        assert not mobile.input_context(OWNER)[0]['attachment_ids']
        assert not list(files.root().iterdir())
    asyncio.run(run())


def test_qq_file_resolves_id_never_trusts_event_url_or_local_path():
    async def run():
        incoming = event([{'type':'file','data':{'file_id':'opaque-id','name':'note.txt','url':'http://localhost/secret','path':'C:/secret'}}])
        api = AsyncMock(return_value={'data':{'url':'https://multimedia.nt.qq.com/download'}})
        download = AsyncMock(return_value=b'synthetic note')
        with patch.object(onebot_transport_service, 'request_task_api', api), patch.object(files, 'download', download):
            await onebot.process_onebot_event(incoming, AsyncMock())
        assert api.await_args.args[1:] == ('get_private_file_url', {'file_id':'opaque-id'})
        download.assert_awaited_once_with('https://multimedia.nt.qq.com/download')
        assert len(mobile.input_context(OWNER)[0]['attachment_ids']) == 1
    asyncio.run(run())


def test_mobile_events_bypass_debounce_and_group_never_has_mobile_origin():
    assert not onebot.OneBotEventDebouncer(AsyncMock()).enqueue(event('/任务 暂停'))
    assert mobile.origin_from_event({**event(), 'message_type':'group', 'group_id':'42'}) is None
    assert mobile.origin_from_event(event(user='99999')) is None


def test_followup_revises_same_task_and_cancel_cannot_resume():
    async def run():
        task = handoff()
        result = await mobile.amend(OWNER, task['id'], '控制在两页以内，保留原文。')
        updated = tasks.get(task['id'])
        assert result['task_id'] == task['id']
        assert updated['revision'] > task['revision']
        assert updated['snapshot']['updates'][-1] == '控制在两页以内，保留原文。'
        assert updated['status'] == 'ready'
        await mobile.handle_command({}, OWNER, '/任务 暂停', 'pause')
        assert tasks.get(task['id'])['status'] == 'paused'
        await mobile.handle_command({}, OWNER, '/任务 继续', 'resume')
        assert tasks.get(task['id'])['status'] == 'ready'
        await mobile.handle_command({}, OWNER, '/任务 取消', 'cancel')
        assert tasks.get(task['id'])['status'] == 'cancelled'
        with pytest.raises(ValueError): await mobile.amend(OWNER, task['id'], '继续')
    asyncio.run(run())


def test_durable_outbox_waits_for_correct_bot_and_stores_ack():
    async def run():
        handoff()
        api = AsyncMock(return_value={'status':'ok','retcode':0,'data':{'message_id':123}})
        with patch.object(onebot_transport_service, 'request_task_api', api):
            assert await mobile.flush_outbox() == 0
            api.assert_not_awaited()
            with patch.object(onebot_transport_service, 'has_task_connection', return_value=True):
                assert await mobile.flush_outbox() == 1
                assert await mobile.flush_outbox() == 0
        assert outbox()[0]['status'] == 'sent'
        assert json.loads(outbox()[0]['receipt_json'])['message_id'] == 123
    asyncio.run(run())


@pytest.mark.parametrize('error,state', [(TimeoutError(),'unknown'), (onebot.TaskApiRejected('拒绝'),'failed')])
def test_uncertain_or_rejected_sends_not_automatically_repeated(error, state):
    async def run():
        handoff()
        api = AsyncMock(side_effect=error)
        with patch.object(onebot_transport_service, 'has_task_connection', return_value=True), patch.object(onebot_transport_service, 'request_task_api', api):
            await mobile.flush_outbox(); await mobile.flush_outbox()
        assert api.await_count == 1
        assert outbox()[0]['status'] == state
    asyncio.run(run())


def test_restart_recovers_inflight_as_unknown_without_resend():
    handoff()
    with closing(db.get_conn()) as conn, conn:
        conn.execute("UPDATE mobile_task_outbox SET status='sending'")
    mobile.recover_delivery()
    assert outbox()[0]['status'] == 'unknown'


def test_final_report_waits_for_matching_run_and_attaches_real_artifact():
    task = handoff()
    db.save_message('assistant', '旧结果', conversation_id=task['conversation_id'], request_id='old')
    tasks.update(task['id'], status='completed')
    mobile.report_task(tasks.get(task['id']))
    assert len(outbox()) == 1
    with closing(db.get_conn()) as conn:
        request = conn.execute('SELECT request_id FROM agent_task_runs WHERE task_id=?', (task['id'],)).fetchone()[0]
    db.save_message('assistant','新结果',conversation_id=task['conversation_id'],request_id=request)
    artifact = agent_file_service.write_document('报告.docx', '真实测试报告', task['id'])
    mobile.report_task(tasks.get(task['id'])); mobile.report_task(tasks.get(task['id']))
    assert len(outbox()) == 3
    assert any('新结果' in row['payload_json'] for row in outbox())
    assert not any('旧结果' in row['payload_json'] for row in outbox())
    assert any('artifact_id' in row['payload_json'] for row in outbox())


def test_document_task_executes_real_tools_and_returns_verified_file():
    async def run():
        await files.store_document(OWNER, 'input', '材料.txt', '测试材料：营收 10 万。'.encode())
        task = handoff()
        with closing(db.get_conn()) as conn:
            path = conn.execute('SELECT path FROM mobile_task_attachments').fetchone()[0]
        call = lambda tool, **args: ToolCall(tool,tool,json.dumps(args,ensure_ascii=False))
        planner = AsyncMock(side_effect=[_completion(call('agent_read_document', path=path)),
            _completion(call('agent_write_document', name='摘要.docx', content='营收为 10 万。')), _completion()])
        with patch('app.agent_loop_service.call_chat_completion_result', planner), \
             patch('app.llm.call_chat_completion_result', AsyncMock(return_value=replace(_completion(), content='报告已完成。'))):
            await tasks.continue_task(task['id'])
        current = tasks.get(task['id'])
        assert current['status'] == 'completed', (current['snapshot'], tasks.observations(task['id']))
        assert all(item['status'] == 'completed' for item in tasks.observations(task['id']))
        artifacts = mobile.artifacts_for(current)
        assert len(artifacts) == 1
        file, _ = artifact_service.file(artifacts[0]['id'])
        assert '10 万' in '\n'.join(p.text for p in Document(file).paragraphs)
        mobile.report_task(current)
        api = AsyncMock(return_value={'status':'ok','retcode':0})
        with patch.object(onebot_transport_service, 'has_task_connection',return_value=True), patch.object(onebot_transport_service,'request_task_api',api):
            await mobile.flush_outbox()
        assert any(call.args[1] == 'upload_private_file' for call in api.await_args_list)
    asyncio.run(run())
