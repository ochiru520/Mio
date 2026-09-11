"""User-authorized document reads and versioned outputs for Agent tasks."""

from __future__ import annotations

import hashlib
import io
import json
import mimetypes
import os
import re
import zipfile
from pathlib import Path
from urllib.parse import quote

from docx import Document
import pypdfium2

from . import db
from .config import settings

TEXT = {".txt", ".md", ".json", ".csv", ".tsv", ".yaml", ".yml"}
READABLE = TEXT | {".docx", ".pdf", ".png", ".jpg", ".jpeg", ".webp", ".mp4", ".webm", ".mov"}


def workspace() -> Path:
    root = settings.data_dir.resolve() / "Agent文件"
    root.mkdir(parents=True, exist_ok=True)
    if root.resolve() != root:
        raise ValueError("Agent 成果目录不能是指向其他位置的联接或符号链接。")
    return root.resolve()


def _private_path(path: Path, *, grant: bool = False) -> bool:
    data = settings.data_dir.resolve()
    if path == data or path.is_relative_to(data) or (grant and data.is_relative_to(path)):
        return True
    # Some installations keep their settings outside the main data directory.
    sensitive = [settings.db_path, settings.runtime_config_path, settings.model_profiles_path,
                 settings.companion_config_path, settings.mio_profile_path]
    return any(path == item.resolve() or (grant and item.resolve().is_relative_to(path)) for item in sensitive)


def roots(changes: list[str] | None = None) -> list[str]:
    from .agent_task_service import initialize
    initialize()
    with db.get_conn() as conn:
        conn.execute("CREATE TABLE IF NOT EXISTS agent_file_roots(path TEXT PRIMARY KEY)")
        if changes is not None:
            resolved = []
            for item in changes:
                path = Path(item).resolve(strict=True)
                if not path.is_dir() or path.parent == path or _private_path(path, grant=True):
                    raise ValueError("请选择具体资料目录，不能授权整个磁盘或 Mio 私人数据根目录。")
                resolved.append(str(path))
            conn.execute("DELETE FROM agent_file_roots")
            conn.executemany("INSERT OR IGNORE INTO agent_file_roots VALUES(?)", [(p,) for p in resolved])
        values = [row[0] for row in conn.execute("SELECT path FROM agent_file_roots ORDER BY path")]
    values = [value for value in values if not _private_path(Path(value).resolve(), grant=True)]
    return list(dict.fromkeys([str(workspace()), *values]))


def resolve_read(value: str) -> Path:
    path = Path(value)
    path = (path if path.is_absolute() else workspace() / path).resolve(strict=True)
    if _private_path(path) and not path.is_relative_to(workspace()):
        raise ValueError("私人数据不能通过文件工具读取，请使用专用的记忆或日记工具。")
    if not any(path.is_relative_to(Path(root)) for root in roots()):
        raise ValueError("该路径尚未授权，请在 Agent 设置的文件权限中添加资料目录。")
    if any(part.startswith('.') for part in path.parts) or ':' in path.name:
        raise ValueError("不允许读取隐藏配置或命名数据流。")
    if path.is_file() and path.suffix.lower() not in READABLE:
        raise ValueError("不支持该文件类型。")
    return path


def list_files(path: str = "", limit: int = 100) -> dict:
    if not path:
        return {"roots": roots(), "output_root": str(workspace())}
    directory = resolve_read(path)
    if not directory.is_dir():
        raise ValueError("需要目录路径。")
    items = []
    for item in sorted(directory.iterdir(), key=lambda p: (not p.is_dir(), p.name.casefold())):
        if item.name.startswith('.') or (item.is_file() and item.suffix.lower() not in READABLE):
            continue
        try:
            target = resolve_read(str(item))
        except (ValueError, OSError):
            continue
        items.append({"path": str(target), "name": item.name, "directory": target.is_dir(),
                      "size": target.stat().st_size if target.is_file() else 0})
        if len(items) >= limit:
            break
    return {"files": items, "root": str(directory)}


def read_document(path: str, offset: int = 0, max_chars: int = 16000) -> dict:
    source = resolve_read(path)
    if not source.is_file() or source.stat().st_size > 32 * 1024 * 1024:
        raise ValueError("文档必须存在且小于 32MB。")
    suffix = source.suffix.lower()
    if suffix in TEXT:
        content = source.read_text(encoding="utf-8-sig")
    elif suffix == ".docx":
        with zipfile.ZipFile(source) as archive:
            if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
                raise ValueError("文档解压体积过大。")
        doc = Document(source)
        paragraphs = [p.text for p in doc.paragraphs]
        paragraphs.extend("\t".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
        content = "\n".join(paragraphs)
    elif suffix == ".pdf":
        with pypdfium2.PdfDocument(source) as pdf:
            parts = []
            for index in range(min(len(pdf), 200)):
                with pdf[index] as page:
                    with page.get_textpage() as textpage:
                        parts.append(f"[第 {index + 1} 页]\n{textpage.get_text_range()}")
            content = "\n".join(parts)
    else:
        raise ValueError("文档读取支持 TXT、Markdown、JSON、CSV、YAML、DOCX 和有文本层的 PDF。")
    return {"path": str(source), "text": content[offset:offset + max_chars],
            "offset": offset, "total_chars": len(content), "has_more": offset + max_chars < len(content),
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest()}


def write_document(name: str, content: str, task_id: str) -> dict:
    filename = Path(name).name
    if filename != name or ':' in name or name.startswith('.') or not re.fullmatch(r"task_[a-zA-Z0-9_-]+", task_id):
        raise ValueError("输出需要普通文件名与有效任务。")
    suffix = Path(name).suffix.lower()
    if suffix not in TEXT | {".docx"}:
        raise ValueError("输出支持 TXT、Markdown、JSON、CSV、YAML 和 DOCX。")
    directory = workspace() / task_id
    directory.mkdir(exist_ok=True)
    if directory.resolve() != directory:
        raise ValueError("任务成果目录不能是联接或符号链接。")
    if suffix == ".json":
        json.loads(content)
    if suffix == ".docx":
        doc = Document()
        for line in content.splitlines():
            doc.add_paragraph(line)
        buffer = io.BytesIO()
        doc.save(buffer)
        payload = buffer.getvalue()
    else:
        payload = content.encode("utf-8")
    digest = hashlib.sha256(payload).hexdigest()
    output = directory / f"{Path(name).stem}-{digest[:10]}{suffix}"
    if not output.exists():
        with output.open('xb') as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
    verified = read_document(str(output), max_chars=300)
    if verified["sha256"] != digest:
        raise ValueError("输出文件内容校验失败。")
    from . import artifact_service
    artifact = artifact_service.import_file(output, source_root=directory, producer="agent_document",
        expected_sha256=digest, task_id=task_id, name=name)
    return {"path": str(output), "name": output.name, "size": output.stat().st_size,
            "sha256": digest, "verified": True, "preview": verified["text"],
            "artifact_id": artifact["id"], "artifact_url": artifact["url"], "version": artifact["version"],
            "url": f"/api/agent/work/files/{task_id}/{quote(output.name, safe='')}"}


def output_file(task_id: str, name: str) -> Path:
    if not re.fullmatch(r"task_[a-zA-Z0-9_-]+", task_id) or Path(name).name != name:
        raise ValueError("无效输出路径。")
    candidate = (workspace() / task_id / name).resolve(strict=True)
    if not candidate.is_relative_to(workspace() / task_id) or not candidate.is_file() or candidate.suffix.lower() not in READABLE:
        raise ValueError("找不到输出文件。")
    return candidate
