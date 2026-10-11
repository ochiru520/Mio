from contextlib import closing
import io
import hashlib
import json
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from openpyxl import load_workbook
import pypdfium2
import pytest

from app import agent_file_service as files, artifact_service as artifacts
from app import office_document_service as office
from test_runtime_safety_v2 import isolated


def test_rich_docx_preserves_headings_tables_lists_and_source_links(isolated):
    content='# 合成验收报告\n\n**关键结论**：存在原始证据。\n- 第一项\n- 第二项\n\n| 指标 | 数值 |\n| --- | --- |\n| 合成指标 | 42 |\n\n[官方来源](https://example.com/source)'
    output=files.write_document('report.docx',content,'task_office')
    document=Document(output['path'])
    assert document.paragraphs[0].style.name=='Heading 1'
    assert document.tables[0].cell(1,1).text=='42'
    assert any(p.style.name=='List Bullet' for p in document.paragraphs)
    assert any(rel.target_ref=='https://example.com/source' for rel in document.part.rels.values())
    assert output['quality_status']=='not_reviewed' and output['verification_scope']=='file_integrity_and_structure'
    assert artifacts.get(output['artifact_id'])['integrity_status']=='verified'


def test_xlsx_retains_formulas_structure_and_versions_without_claiming_recalculation(isolated):
    spec={'sheets':[{'name':'数据','rows':[['项目','数值'],['合成',42],['公式','=SUM(B2:B2)']], 'widths':[24,18]},
                    {'name':'来源','rows':[['日期','说明'],['2026-10-04','合成验收']]}]}
    output=files.write_document('report.xlsx',json.dumps(spec,ensure_ascii=False),'task_office')
    book=load_workbook(output['path'],keep_links=False)
    try:
        assert book.sheetnames==['数据','来源'] and book['数据']['B3'].value=='=SUM(B2:B2)'
        assert book['数据'].freeze_panes=='A2' and book['数据']['A1'].font.bold
    finally: book.close()
    assert output['formulas']==1 and not output['formula_values_recalculated']
    again=files.write_document('report.xlsx',json.dumps(spec,ensure_ascii=False),'task_office')
    # OOXML ZIP timestamps may change bytes between invocations; both outputs remain immutable.
    assert files.read_document(again['path'])['text']==files.read_document(output['path'])['text']
    for formula in ('=WEBSERVICE("https://example.com")','=cmd|anything!A1',"='[outside.xlsx]Sheet1'!A1"):
        with pytest.raises(ValueError,match='外部资源'):
            files.write_document('unsafe.xlsx',json.dumps({'sheets':[{'name':'数据','rows':[[formula]]}]}),'task_office')


def test_cjk_pdf_is_readable_renders_multiple_pages_and_archives_verified_bytes(isolated):
    content='# 合成中文报告\n\n**说明**：这是程序输出验收。\n\n| 名称 | 状态 |\n| --- | --- |\n| 合成资料 | 已核对 |\n\n'+('\n'.join(f'第 {i} 项：来源日期 2026-10-04，保留原始证据。' for i in range(90)))
    output=files.write_document('report.pdf',content,'task_office')
    result=files.read_document(output['path'],max_chars=200000)
    assert '合成中文报告' in result['text'] and '第 89 项' in result['text']
    with closing(pypdfium2.PdfDocument(output['path'])) as pdf:
        assert len(pdf)>=3
        for index in (0,len(pdf)-1):
            with closing(pdf[index]) as page, closing(page.render(scale=1)) as bitmap:
                image=bitmap.to_pil()
                assert image.convert('L').getextrema()[0]<100
    blob,artifact=artifacts.file(output['artifact_id'])
    assert hashlib.sha256(blob.read_bytes()).hexdigest()==output['sha256'] and artifact['mime_type']=='application/pdf'


@pytest.mark.parametrize('cache', ['missing', 'error', 'number', 'zero', 'boolean', 'text', 'changed_formula', 'partial'])
def test_recalculation_checks_formula_caches_not_numeric_inputs(cache):
    """Explicit converter stand-in; proves reporting, not a native engine."""
    spec = {'sheets': [{'name': '数据', 'rows': [['值', '公式'], [42, '=A2*2'], [7, '=A3*2']]}]}
    with patch.object(office.shutil, 'which', return_value=None):
        original, metadata = office.xlsx_bytes(json.dumps(spec, ensure_ascii=False))
    assert not metadata['formula_values_recalculated']

    def convert(command, **kwargs):
        assert command[1].startswith('-env:UserInstallation=file:///')
        assert kwargs['timeout'] == 45
        source = Path(command[-1])
        output = Path(command[command.index('--outdir') + 1]) / source.name
        with ZipFile(io.BytesIO(source.read_bytes())) as zin, ZipFile(output, 'w') as zout:
            for entry in zin.infolist():
                blob = zin.read(entry.filename)
                if entry.filename == 'xl/worksheets/sheet1.xml':
                    root = ET.fromstring(blob)
                    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                    for cell in root.iter(ns + 'c'):
                        expression = cell.find(ns + 'f')
                        if expression is None:
                            continue
                        value = cell.find(ns + 'v')
                        if value is None:
                            value = ET.SubElement(cell, ns + 'v')
                        if cache == 'missing' or cache == 'partial' and cell.get('r') == 'B3':
                            continue
                        if cache == 'error':
                            cell.set('t', 'e'); value.text = '#VALUE!'
                        elif cache == 'boolean':
                            cell.set('t', 'b'); value.text = '0'
                        elif cache == 'text':
                            cell.set('t', 'str'); value.text = 'synthetic'
                        else:
                            cell.set('t', 'n'); value.text = '0' if cache == 'zero' else '84' if cell.get('r') == 'B2' else '14'
                        if cache == 'changed_formula':
                            expression.text = 'A2*3'
                    blob = ET.tostring(root, encoding='utf-8')
                zout.writestr(entry, blob)

    with patch.object(office.shutil, 'which', return_value='synthetic-soffice.exe'), \
         patch.object(office.subprocess, 'run', side_effect=convert) as run:
        result, metadata = office.xlsx_bytes(json.dumps(spec, ensure_ascii=False))
    run.assert_called_once()
    expected = cache in {'number', 'zero', 'boolean', 'text'}
    assert metadata['formula_values_recalculated'] == expected
    assert metadata['recalculation_engine'] == ('LibreOffice' if expected else 'none')
    with closing(load_workbook(io.BytesIO(result), read_only=True)) as book:
        assert book['数据']['B2'].value == '=A2*2' and book['数据']['B3'].value == '=A3*2'
        assert book['数据']['A2'].value == 42
    with closing(load_workbook(io.BytesIO(result), data_only=True, read_only=True)) as book:
        if not expected:
            assert book['数据']['B2'].value is None
        elif cache == 'zero':
            assert book['数据']['B2'].value == 0
        elif cache == 'boolean':
            assert book['数据']['B2'].value is False


@pytest.mark.parametrize('failure', ['timeout', 'missing_output', 'corrupt_output'])
def test_recalculation_failure_preserves_original_formulas(failure):
    import subprocess
    spec = {'sheets': [{'name': '数据', 'rows': [[42, '=A1*2']]}]}
    def corrupt(command, **kwargs):
        source = Path(command[-1])
        output = Path(command[command.index('--outdir') + 1]) / source.name
        output.write_bytes(b'corrupt synthetic converter output')
    effect = subprocess.TimeoutExpired('synthetic-soffice', 45) if failure == 'timeout' else corrupt if failure == 'corrupt_output' else None
    with patch.object(office.shutil, 'which', return_value='synthetic-soffice.exe'), \
         patch.object(office.subprocess, 'run', side_effect=effect):
        result, metadata = office.xlsx_bytes(json.dumps(spec))
    assert not metadata['formula_values_recalculated'] and metadata['recalculation_engine'] == 'none'
    with closing(load_workbook(io.BytesIO(result), read_only=True)) as book:
        assert book['数据']['B1'].value == '=A1*2'
