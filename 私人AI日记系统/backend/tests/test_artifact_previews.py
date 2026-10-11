"""Real archived Office files and passive media previews, with isolated user data."""
import hashlib
import io
from zipfile import ZipFile

import pytest
from docx import Document
from docx.shared import Inches
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image

from app import artifact_service as artifacts, artifact_preview_service as preview
from app.routes.artifacts import router
from test_runtime_safety_v2 import isolated


@pytest.fixture
def client():
    app = FastAPI()
    app.include_router(router)
    with TestClient(app) as value:
        yield value


def archive(tmp_path, name, payload, mime_type=''):
    source = tmp_path / name
    source.write_bytes(payload)
    item = artifacts.import_file(source, source_root=tmp_path, producer='preview_test',
                                 expected_sha256=hashlib.sha256(payload).hexdigest(), mime_type=mime_type)
    source.unlink()
    return item


def package(parts):
    buffer = io.BytesIO()
    with ZipFile(buffer, 'w') as value:
        for name, content in parts.items():
            value.writestr(name, content)
    return buffer.getvalue()


def test_docx_preview_reads_verified_blob_preserves_body_order_and_embedded_image(tmp_path, client):
    doc = Document()
    doc.add_heading('中文报告', level=2)
    doc.add_paragraph('首段 <script>仅文本</script>')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '列一'
    table.cell(0, 1).text = '42'
    doc.add_paragraph('表格后的结论')
    image = io.BytesIO()
    Image.new('RGB', (20, 10), 'red').save(image, format='PNG')
    image.seek(0)
    doc.add_picture(image, width=Inches(1))
    buffer = io.BytesIO()
    doc.save(buffer)
    item = archive(tmp_path, '报告.DOCX', buffer.getvalue(), 'application/octet-stream')
    response = client.get(item['preview_url'])
    assert response.status_code == 200
    body = response.json()
    assert body['kind'] == 'document' and body['format'] == 'docx' and body['supported']
    assert body['blocks'][:4] == [
        {'type': 'heading', 'text': '中文报告', 'level': 2},
        {'type': 'paragraph', 'text': '首段 <script>仅文本</script>'},
        {'type': 'table', 'rows': [['列一', '42']]},
        {'type': 'paragraph', 'text': '表格后的结论'},
    ]
    assert body['blocks'][4]['url'].startswith('data:image/png;base64,')
    assert not body['truncated'] and '原始版式' in body['notice']
    assert client.get(item['url']).content == buffer.getvalue()


def test_workbook_preview_shows_multiple_sheets_formulas_and_sparse_cells_without_calculating(tmp_path, client):
    book = Workbook()
    sheet = book.active
    sheet.title = '数据'
    sheet.append(['项目', '值'])
    sheet.append(['甲', 42])
    sheet.append(['公式', '=B2*2'])
    sheet['D5'] = True
    book.create_sheet('空表')
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    item = archive(tmp_path, '数据.xlsx', buffer.getvalue())
    response = client.get(item['preview_url'])
    assert response.status_code == 200
    body = response.json()
    assert body['sheets'][0]['rows'] == [['项目', '值'], ['甲', '42'], ['公式', '=B2*2'], ['', '', '', 'TRUE']]
    assert body['sheets'][0]['row_numbers'] == [1, 2, 3, 5]
    assert body['sheets'][1]['rows'] == []
    assert '不执行重算' in body['notice'] and not body['truncated']


def test_workbook_preview_reports_row_column_and_character_truncation(tmp_path, client, monkeypatch):
    book = Workbook()
    for row in range(201):
        book.active.append([f'行{row}'] + [''] * 60 + ['超列'])
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    item = archive(tmp_path, '大表.xlsx', buffer.getvalue())
    body = client.get(item['preview_url']).json()
    assert len(body['sheets'][0]['rows']) == 200
    assert body['sheets'][0]['truncated'] and body['truncated']
    assert all(len(row) <= 60 for row in body['sheets'][0]['rows'])
    assert '已截断' in body['notice']
    monkeypatch.setattr(preview, 'MAX_TEXT_CHARS', 6)
    body = client.get(item['preview_url']).json()
    assert body['truncated']
    assert sum(len(cell) for row in body['sheets'][0]['rows'] for cell in row) <= 6


def test_presentation_preview_uses_recorded_slide_order_and_tables(tmp_path, client):
    p, a, r = preview.NS['p'], preview.NS['a'], preview.NS['r']
    presentation = f'<p:presentation xmlns:p="{p}" xmlns:r="{r}"><p:sldIdLst><p:sldId id="256" r:id="second"/><p:sldId id="257" r:id="first"/></p:sldIdLst></p:presentation>'
    relations = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="first" Target="slides/slide1.xml"/><Relationship Id="second" Target="slides/slide2.xml"/></Relationships>'
    def slide(text):
        return f'<p:sld xmlns:p="{p}" xmlns:a="{a}"><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp><p:graphicFrame><a:graphic><a:graphicData><a:tbl><a:tr><a:tc><a:txBody><a:p><a:r><a:t>表格</a:t></a:r></a:p></a:txBody></a:tc></a:tr></a:tbl></a:graphicData></a:graphic></p:graphicFrame></p:spTree></p:cSld></p:sld>'
    payload = package({'ppt/presentation.xml': presentation, 'ppt/_rels/presentation.xml.rels': relations,
                       'ppt/slides/slide1.xml': slide('后显示'), 'ppt/slides/slide2.xml': slide('先显示')})
    item = archive(tmp_path, '讲义.pptx', payload)
    body = client.get(item['preview_url']).json()
    assert body['format'] == 'pptx'
    assert [value['blocks'][0]['text'] for value in body['slides']] == ['先显示', '后显示']
    assert body['slides'][0]['blocks'][1] == {'type': 'table', 'rows': [['表格']]}


@pytest.mark.parametrize('encoding', ['utf-8', 'utf-16'])
def test_office_preview_rejects_xml_entities_without_loading_external_sources(tmp_path, client, encoding):
    w = preview.NS['w']
    xml = f'<?xml version="1.0" encoding="{encoding}"?><!DOCTYPE document [<!ENTITY secret SYSTEM "file:///C:/Windows/win.ini">]><w:document xmlns:w="{w}"><w:body><w:p><w:r><w:t>&secret;</w:t></w:r></w:p></w:body></w:document>'
    item = archive(tmp_path, 'unsafe.docx', package({'word/document.xml': xml.encode(encoding)}))
    response = client.get(item['preview_url'])
    assert response.status_code == 422 and '不安全 XML' in response.json()['detail']


def test_corrupt_large_and_tampered_office_files_produce_actionable_errors(tmp_path, client, monkeypatch):
    corrupt = archive(tmp_path, '损坏.docx', b'not a zip')
    response = client.get(corrupt['preview_url'])
    assert response.status_code == 422 and '损坏' in response.json()['detail']
    monkeypatch.setattr(preview, 'MAX_FILE_BYTES', 4)
    assert '32MB' in client.get(corrupt['preview_url']).json()['detail']
    blob, _ = artifacts.file(corrupt['id'])
    blob.write_bytes(b'tampered')
    response = client.get(corrupt['preview_url'])
    assert response.status_code == 404 and '校验失败' in response.json()['detail']


def test_archive_expansion_limit_and_image_budget_are_reported(tmp_path, client, monkeypatch):
    doc = Document()
    doc.add_paragraph('正文')
    image = io.BytesIO()
    Image.new('RGB', (20, 20)).save(image, format='PNG')
    image.seek(0)
    doc.add_picture(image)
    buffer = io.BytesIO()
    doc.save(buffer)
    item = archive(tmp_path, '图片.docx', buffer.getvalue())
    monkeypatch.setattr(preview, 'MAX_IMAGE_BYTES', 4)
    body = client.get(item['preview_url']).json()
    assert body['blocks'] == [{'type': 'paragraph', 'text': '正文'}] and body['truncated']
    monkeypatch.setattr(preview, 'MAX_EXPANDED_BYTES', 4)
    response = client.get(item['preview_url'])
    assert response.status_code == 422 and '解压体积' in response.json()['detail']


def test_generic_image_mime_gets_inline_content_and_svg_is_passive_text(tmp_path, client):
    image = io.BytesIO()
    Image.new('RGB', (20, 10)).save(image, format='PNG')
    item = archive(tmp_path, '图片.png', image.getvalue(), 'application/octet-stream')
    body = client.get(item['preview_url']).json()
    assert body['kind'] == 'image'
    response = client.get(body['url'])
    assert response.content == image.getvalue() and response.headers['content-type'] == 'image/png'
    assert response.headers['content-disposition'].startswith('inline')
    svg = archive(tmp_path, '图.svg', b'<svg><script>danger()</script></svg>', 'image/svg+xml')
    assert client.get(svg['preview_url']).json()['kind'] == 'text'
    assert client.get('/api/artifacts/' + svg['id'] + '/content').headers['content-disposition'].startswith('attachment')


def test_text_truncation_counts_characters_and_existing_pdf_preview_remains_inline(tmp_path, client):
    item = archive(tmp_path, '中文.txt', ('中' * 40000).encode('utf-8'))
    body = client.get(item['preview_url']).json()
    assert len(body['content']) == 40000 and not body['truncated']
    pdf = archive(tmp_path, '文件.pdf', b'%PDF-1.4\n%%EOF')
    body = client.get(pdf['preview_url']).json()
    assert body['kind'] == 'pdf' and body['url'].endswith('/content')
