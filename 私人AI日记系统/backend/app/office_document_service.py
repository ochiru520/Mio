"""Local office output builders. Structural verification is separate from quality review."""
from contextlib import closing
import html
import io
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError


def _formula_cache_verified(original, converted):
    """Check every formula's retained expression and actual cached result.

    Ordinary numeric input cells do not establish calculation success. This
    verifies the converter's saved caches, not business correctness or layout.
    """
    from openpyxl import load_workbook
    with closing(load_workbook(io.BytesIO(original), read_only=True, keep_links=False)) as before, \
         closing(load_workbook(io.BytesIO(converted), read_only=True, keep_links=False)) as after, \
         closing(load_workbook(io.BytesIO(converted), data_only=True, read_only=True, keep_links=False)) as values:
        if before.sheetnames != after.sheetnames or before.sheetnames != values.sheetnames:
            return False
        formula_count = 0
        for sheet in before.worksheets:
            target, cached = after[sheet.title], values[sheet.title]
            # Stream each XML sheet once. Repeated read-only coordinate lookup
            # reparses a sheet and becomes quadratic for larger workbooks.
            source_cells = {cell.coordinate: (cell.value, cell.data_type)
                            for row in sheet.iter_rows() for cell in row if cell.value is not None}
            target_cells = {cell.coordinate: (cell.value, cell.data_type)
                            for row in target.iter_rows() for cell in row if cell.value is not None}
            if source_cells != target_cells:
                return False
            formula_cells = {coordinate for coordinate, (_, kind) in source_cells.items() if kind == 'f'}
            results = {cell.coordinate: (cell.value, cell.data_type)
                       for row in cached.iter_rows() for cell in row
                       if getattr(cell, 'coordinate', None) in formula_cells}
            formula_count += len(formula_cells)
            for coordinate in formula_cells:
                value, kind = results.get(coordinate, (None, 'n'))
                if value is None or kind == 'e':
                    return False
                if isinstance(value, (float, int)) and not math.isfinite(value):
                    return False
        return formula_count > 0


def blocks(markdown):
    lines = markdown.splitlines(); index = 0; fenced = False
    while index < len(lines):
        line = lines[index]; index += 1
        if line.startswith('```'):
            fenced = not fenced; continue
        if fenced:
            yield 'code', line; continue
        if '|' in line and index < len(lines) and re.fullmatch(r'\s*\|?\s*:?-+:?\s*(?:\|\s*:?-+:?\s*)+\|?\s*', lines[index]):
            rows = [[cell.strip() for cell in line.strip().strip('|').split('|')]]
            index += 1
            while index < len(lines) and '|' in lines[index] and lines[index].strip():
                rows.append([cell.strip() for cell in lines[index].strip().strip('|').split('|')]); index += 1
            width = max(map(len, rows))
            yield 'table', [row+['']*(width-len(row)) for row in rows]
        elif re.match(r'^#{1,6}\s+',line):
            level, text = line.split(' ',1); yield 'heading', (len(level),text)
        elif re.match(r'^\s*[-*+]\s+',line):
            yield 'bullet', re.sub(r'^\s*[-*+]\s+','',line)
        elif re.match(r'^\s*\d+[.)]\s+',line):
            yield 'number', line.strip()
        else:
            yield 'paragraph',line


def inline_docx(paragraph, text):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    tokens = re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\))',text)
    for token in tokens:
        link=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',token)
        if link:
            element=OxmlElement('w:hyperlink'); element.set(qn('r:id'),paragraph.part.relate_to(link[2],RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
            run=OxmlElement('w:r'); properties=OxmlElement('w:rPr'); style=OxmlElement('w:rStyle'); style.set(qn('w:val'),'Hyperlink'); properties.append(style); run.append(properties)
            value=OxmlElement('w:t'); value.text=link[1]; run.append(value); element.append(run); paragraph._p.append(element)
        else:
            run=paragraph.add_run(token[2:-2] if token.startswith('**') else token[1:-1] if token.startswith('`') else token)
            if token.startswith('**'): run.bold=True
            if token.startswith('`'): run.font.name='Consolas'


def docx_bytes(content):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml.ns import qn
    doc=Document(); section=doc.sections[0]
    section.top_margin=section.bottom_margin=Cm(2); section.left_margin=section.right_margin=Cm(2.2)
    normal=doc.styles['Normal']; normal.font.name='Microsoft YaHei'; normal.font.size=Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    normal.paragraph_format.space_after=Pt(6)
    for kind,value in blocks(content):
        if kind=='table':
            table=doc.add_table(rows=len(value),cols=len(value[0])); table.style='Table Grid'
            for row_index,row in enumerate(value):
                for column,text in enumerate(row):
                    paragraph=table.cell(row_index,column).paragraphs[0]; inline_docx(paragraph,text)
                    if row_index==0:
                        for run in paragraph.runs: run.bold=True
        else:
            if kind=='heading': paragraph=doc.add_heading(level=value[0]); value=value[1]
            else: paragraph=doc.add_paragraph(style='List Bullet' if kind=='bullet' else 'List Number' if kind=='number' else None)
            if kind=='number': value=re.sub(r'^\d+[.)]\s+','',value)
            inline_docx(paragraph,value)
            if kind=='code':
                for run in paragraph.runs: run.font.name='Consolas'
    buffer=io.BytesIO(); doc.save(buffer)
    Document(io.BytesIO(buffer.getvalue()))
    return buffer.getvalue()


def xlsx_bytes(content):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    spec=json.loads(content)
    if not isinstance(spec,dict) or set(spec)-{'sheets'} or not isinstance(spec.get('sheets'),list) or not 1<=len(spec['sheets'])<=30:
        raise ValueError('XLSX 正文需要 JSON：{"sheets":[{"name":"表名","rows":[["标题",1],["公式","=B1*2"]]}]}。最多 30 张表。')
    book=Workbook(); book.remove(book.active); total=0; formulas=0; names=set()
    for sheet in spec['sheets']:
        if not isinstance(sheet,dict) or set(sheet)-{'name','rows','widths'}:
            raise ValueError('工作表只支持 name、rows、widths。')
        name=str(sheet.get('name') or '').strip()
        if not name or len(name)>31 or re.search(r'[\\/*?:\[\]]',name) or name.lower() in names:
            raise ValueError('工作表名称无效或重复。')
        names.add(name.lower()); rows=sheet.get('rows')
        if not isinstance(rows,list) or len(rows)>20000:
            raise ValueError('工作表 rows 需要是不超过 20000 行的二维数组。')
        ws=book.create_sheet(name)
        for row_index,row in enumerate(rows,1):
            if not isinstance(row,list) or len(row)>100:
                raise ValueError('每一行需要是不超过 100 列的数组。')
            total+=len(row)
            if total>200000: raise ValueError('工作簿超过 20 万单元格，请拆分输出。')
            for column,value in enumerate(row,1):
                if value is not None and (not isinstance(value,(str,int,float,bool)) or isinstance(value,str) and len(value)>32767):
                    raise ValueError('单元格需要文本、数值、布尔值或空值。')
                if isinstance(value,str) and value.startswith('='):
                    if '[' in value or re.search(r'\b(?:WEBSERVICE|HYPERLINK|RTD|DDE|CALL|EXEC|REGISTER|IMPORTXML)\s*\(',value,re.I) or '|' in value:
                        raise ValueError('工作簿公式不能请求外部资源或执行命令。')
                    formulas+=1
                cell=ws.cell(row_index,column,value); cell.alignment=Alignment(vertical='top',wrap_text=True)
                if row_index==1:
                    cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='476C60')
        widths=sheet.get('widths',[])
        if not isinstance(widths,list) or len(widths)>100: raise ValueError('widths 需要列宽数组。')
        for column in range(1,ws.max_column+1):
            width=widths[column-1] if column<=len(widths) else 24
            if not isinstance(width,(float,int)) or not 5<=width<=80: raise ValueError('列宽需要在 5–80 之间。')
            ws.column_dimensions[get_column_letter(column)].width=width
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    buffer=io.BytesIO(); book.calculation.fullCalcOnLoad = True; book.calculation.forceFullCalc = True
    book.save(buffer); book.close()
    output = buffer.getvalue(); recalculated = False
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if formulas and soffice:
        with tempfile.TemporaryDirectory(prefix='mio-xlsx-') as temp:
            source = Path(temp) / 'input.xlsx'; outdir = Path(temp) / 'out'; outdir.mkdir()
            source.write_bytes(output)
            try:
                profile = Path(temp) / 'profile'
                subprocess.run([soffice, '-env:UserInstallation=' + profile.as_uri(), '--headless',
                                '--convert-to', 'xlsx', '--outdir', str(outdir), str(source)],
                               check=True, capture_output=True, timeout=45)
                converted = outdir / source.name
                if converted.is_file():
                    converted_bytes = converted.read_bytes()
                    recalculated = _formula_cache_verified(output, converted_bytes)
                    if recalculated: output = converted_bytes
            except (OSError, subprocess.SubprocessError, ValueError, BadZipFile, ParseError, KeyError):
                recalculated = False
    verified=load_workbook(io.BytesIO(output),read_only=True,keep_links=False); verified.close()
    return output, {'sheets':len(spec['sheets']),'cells':total,'formulas':formulas,'formula_values_recalculated':recalculated,
                    'recalculation_engine': 'LibreOffice' if recalculated else 'none'}


def pdf_font_path():
    configured = os.getenv('MIO_PDF_FONT_PATH', '').strip()
    if configured:
        font = Path(configured).expanduser()
        if not font.is_file():
            raise ValueError('MIO_PDF_FONT_PATH 指定的中文字体不存在。')
        return font
    font_dir = Path(os.getenv('WINDIR', 'C:/Windows')) / 'Fonts'
    for name in ('simhei.ttf', 'msyh.ttc', 'simsun.ttc', 'msjh.ttc'):
        font = font_dir / name
        if font.is_file():
            return font
    raise ValueError('PDF 中文输出需要中文字体，请安装黑体/微软雅黑，或设置 MIO_PDF_FONT_PATH 指向中文 TrueType 字体。')


def pdf_bytes(content):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
    font=pdf_font_path()
    if 'MioCJK' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont('MioCJK',str(font)))
    pdfmetrics.registerFontFamily('MioCJK',normal='MioCJK',bold='MioCJK',italic='MioCJK',boldItalic='MioCJK')
    styles={name:ParagraphStyle(name,fontName='MioCJK',fontSize=size,leading=size*1.6,spaceAfter=8,wordWrap='CJK',alignment=TA_LEFT)
            for name,size in [('body',10),('heading',17),('cell',9)]}
    def text(value):
        escaped=html.escape(str(value))
        escaped=re.sub(r'\*\*([^*]+)\*\*',r'<b>\1</b>',escaped)
        return re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)',r'<link href="\2" color="#476c60">\1</link>',escaped)
    story=[]
    for kind,value in blocks(content):
        if kind=='table':
            rows=[[Paragraph(text(cell),styles['cell']) for cell in row] for row in value]
            table=LongTable(rows,colWidths=[495/len(value[0])]*len(value[0]),repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EDF1EE')),('GRID',(0,0),(-1,-1),.5,colors.HexColor('#CBD4CD')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)])); story.extend([table,Spacer(1,10)])
        else:
            value=value[1] if kind=='heading' else value
            if kind=='bullet': value='- '+value
            story.append(Paragraph(text(value) or ' ',styles['heading' if kind=='heading' else 'body']))
    buffer=io.BytesIO()
    def footer(canvas,document):
        canvas.setFont('MioCJK',8); canvas.drawRightString(A4[0]-50,25,str(document.page))
    SimpleDocTemplate(buffer,pagesize=A4,leftMargin=50,rightMargin=50,topMargin=45,bottomMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)
    import pypdfium2
    with closing(pypdfium2.PdfDocument(buffer.getvalue())) as pdf:
        if not len(pdf): raise ValueError('PDF 输出没有有效页面。')
        for index in range(len(pdf)):
            with closing(pdf[index]) as page, closing(page.render(scale=.5)) as bitmap:
                if bitmap.width<1 or bitmap.height<1: raise ValueError('PDF 页面无法渲染。')
    return buffer.getvalue()
